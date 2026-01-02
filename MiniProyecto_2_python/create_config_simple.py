# create_config_simple.py - Version sin pandas
def create_config_file():
    """Crea config.xlsx manualmente sin usar pandas"""
    
    # Usamos openpyxl directamente
    from openpyxl import Workbook
    
    acme_url = "https://acme-test.uipath.com/login"
    
    # Crear workbook
    wb = Workbook()
    ws = wb.active
    
    # Escribir encabezados
    ws['A1'] = 'acme_url'
    ws['B1'] = 'credential_asset'
    ws['C1'] = 'urgent_threshold'
    
    # Escribir datos
    ws['A2'] = acme_url
    ws['B2'] = 'MiniProyecto_2_Credencial'
    ws['C2'] = 15
    
    # Guardar
    wb.save('config.xlsx')
    
    print("=" * 60)
    print("ARCHIVO CONFIG.XLSX CREADO EXITOSAMENTE")
    print("=" * 60)
    print(f"\nContenido:")
    print(f"URL ACME: {acme_url}")
    print(f"Credencial: MiniProyecto_2_Credencial")
    print(f"Umbral: 15 dias")
    print("\nArchivo guardado como: config.xlsx")

if __name__ == "__main__":
    create_config_file()